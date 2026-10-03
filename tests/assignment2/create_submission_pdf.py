from html import escape
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs' / 'assignment-2' / 'SE3090_Assignment_2_Software_Testing_Report.docx'
OUTPUT = ROOT / 'docs' / 'assignment-2' / 'SE3090_Assignment_2_Software_Testing_Report.pdf'
NAVY = colors.HexColor('#122B5C')
BLUE = colors.HexColor('#1F4E79')
ORANGE = colors.HexColor('#F28C28')
LIGHT_BLUE = colors.HexColor('#EAF2F8')
LIGHT_ORANGE = colors.HexColor('#FFF2E2')
LIGHT_GREY = colors.HexColor('#F3F5F7')
DARK = colors.HexColor('#1F2937')
PAGE_WIDTH, PAGE_HEIGHT = A4
LEFT = 0.62 * inch
RIGHT = 0.62 * inch
TOP = 0.72 * inch
BOTTOM = 0.62 * inch
CONTENT_WIDTH = PAGE_WIDTH - LEFT - RIGHT


def text_of_cell(cell):
    return '<br/>'.join(escape(p.text).replace('\n', '<br/>') for p in cell.paragraphs if p.text) or '&nbsp;'


def paragraph_kind(paragraph):
    style = paragraph.style.name if paragraph.style else ''
    if style.startswith('Heading'):
        return 'heading'
    if style == 'List Bullet':
        return 'bullet'
    if any(run.font.name == 'Consolas' for run in paragraph.runs):
        return 'code'
    return 'body'


def paragraph_text(paragraph):
    return escape(paragraph.text).replace('\n', '<br/>') or '&nbsp;'


def is_page_break(paragraph):
    return 'w:type="page"' in paragraph._p.xml or '<w:br/>' in paragraph._p.xml


def on_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(NAVY)
    canvas.setLineWidth(1.2)
    canvas.line(LEFT, PAGE_HEIGHT - 0.43 * inch, PAGE_WIDTH - RIGHT, PAGE_HEIGHT - 0.43 * inch)
    canvas.setFont('Helvetica-Bold', 8)
    canvas.setFillColor(NAVY)
    canvas.drawString(LEFT, PAGE_HEIGHT - 0.31 * inch, 'SE3090 ASSIGNMENT 2')
    canvas.setFont('Helvetica', 7.5)
    canvas.setFillColor(colors.HexColor('#667085'))
    canvas.drawRightString(PAGE_WIDTH - RIGHT, 0.32 * inch, f'Page {doc.page}')
    canvas.drawString(LEFT, 0.32 * inch, 'Software Testing and Quality Evaluation')
    canvas.restoreState()


def build_styles():
    base = getSampleStyleSheet()
    return {
        'body': ParagraphStyle('body', parent=base['BodyText'], fontName='Helvetica', fontSize=9.2, leading=12, textColor=DARK, spaceAfter=5),
        'bullet': ParagraphStyle('bullet', parent=base['BodyText'], fontName='Helvetica', fontSize=9.1, leading=11.5, leftIndent=14, firstLineIndent=-8, textColor=DARK, spaceAfter=2),
        'heading1': ParagraphStyle('heading1', parent=base['Heading1'], fontName='Helvetica-Bold', fontSize=15, leading=18, textColor=NAVY, spaceBefore=10, spaceAfter=5),
        'heading2': ParagraphStyle('heading2', parent=base['Heading2'], fontName='Helvetica-Bold', fontSize=11.5, leading=14, textColor=BLUE, spaceBefore=7, spaceAfter=4),
        'heading3': ParagraphStyle('heading3', parent=base['Heading3'], fontName='Helvetica-Bold', fontSize=10, leading=12, textColor=BLUE, spaceBefore=5, spaceAfter=3),
        'code': ParagraphStyle('code', parent=base['Code'], fontName='Courier', fontSize=7.3, leading=9, leftIndent=12, textColor=DARK, backColor=LIGHT_GREY, borderPadding=5, spaceAfter=5),
        'cover_small': ParagraphStyle('cover_small', parent=base['BodyText'], fontName='Helvetica', fontSize=11, leading=15, alignment=TA_CENTER, textColor=DARK),
        'cover_title': ParagraphStyle('cover_title', parent=base['Title'], fontName='Helvetica-Bold', fontSize=22, leading=27, alignment=TA_CENTER, textColor=DARK),
    }


def make_table(doc_table, styles):
    rows = []
    for row in doc_table.rows:
        rows.append([Paragraph(text_of_cell(cell), ParagraphStyle('cell', parent=styles['body'], fontSize=7.2, leading=8.8, spaceAfter=0)) for cell in row.cells])
    column_count = len(rows[0]) if rows else 1
    widths = [CONTENT_WIDTH / column_count] * column_count
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), NAVY),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#9AA4B2')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    for row_index in range(1, len(rows)):
        if row_index % 2 == 0:
            table.setStyle(TableStyle([('BACKGROUND', (0, row_index), (-1, row_index), LIGHT_GREY)]))
    return table


def build_story():
    source = Document(SOURCE)
    styles = build_styles()
    story = []
    body = source.element.body
    paragraph_index = 0
    table_index = 0
    seen_cover = False

    for child in body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            paragraph = source.paragraphs[paragraph_index]
            paragraph_index += 1
            text = paragraph.text.strip()
            if not text and not is_page_break(paragraph):
                continue
            if is_page_break(paragraph):
                story.append(PageBreak())
                continue
            if not seen_cover and 'SE3090' in text and len(text) < 30:
                story.append(Spacer(1, 0.65 * inch))
                story.append(Paragraph('SE3090', ParagraphStyle('cover_code', parent=styles['cover_title'], fontSize=28, leading=32, textColor=NAVY)))
                story.append(Paragraph('Software Engineering Frameworks', ParagraphStyle('cover_sub', parent=styles['cover_title'], fontSize=17, leading=21, textColor=BLUE)))
                story.append(Paragraph('Assignment 2', ParagraphStyle('cover_assign', parent=styles['cover_title'], fontSize=30, leading=35, textColor=ORANGE)))
                story.append(Spacer(1, 0.22 * inch))
                seen_cover = True
                continue
            if not seen_cover:
                continue
            kind = paragraph_kind(paragraph)
            if kind == 'heading':
                level = paragraph.style.name.split()[-1]
                style = styles.get('heading' + level, styles['heading1'])
                story.append(Paragraph(paragraph_text(paragraph), style))
            elif kind == 'bullet':
                story.append(Paragraph('&bull;&nbsp; ' + paragraph_text(paragraph), styles['bullet']))
            elif kind == 'code':
                story.append(Paragraph(paragraph_text(paragraph), styles['code']))
            else:
                if text.startswith('Evidence integrity:') or text.startswith('Important interpretation:') or text.startswith('Overall assessment:') or text.startswith('Declaration of evidence status:'):
                    story.append(Table([[Paragraph(paragraph_text(paragraph), styles['body'])]], colWidths=[CONTENT_WIDTH], style=TableStyle([
                        ('BACKGROUND', (0, 0), (-1, -1), LIGHT_ORANGE), ('BOX', (0, 0), (-1, -1), 0.5, ORANGE),
                        ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7),
                        ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ])))
                else:
                    story.append(Paragraph(paragraph_text(paragraph), styles['body']))
        elif tag == 'tbl':
            table = source.tables[table_index]
            table_index += 1
            story.append(KeepTogether([make_table(table, styles), Spacer(1, 5)]))
    return story


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    frame = Frame(LEFT, BOTTOM, CONTENT_WIDTH, PAGE_HEIGHT - TOP - BOTTOM, id='normal', topPadding=5, bottomPadding=5)
    doc = BaseDocTemplate(str(OUTPUT), pagesize=A4, leftMargin=LEFT, rightMargin=RIGHT, topMargin=TOP, bottomMargin=BOTTOM)
    doc.addPageTemplates([PageTemplate(id='main', frames=frame, onPage=on_page)])
    doc.build(build_story())
    print(OUTPUT)


if __name__ == '__main__':
    main()
